"""Controlador para la integración con Microsoft 365 (Microsoft Graph API).

La sesión se inicia una sola vez en el login (model/microsoft_auth.py);
aquí solo se reutiliza, así que no vuelve a pedir cuenta ni contraseña.
"""
import io
import re
import webbrowser

import requests
from docx import Document

from model.microsoft_auth import obtener_token_graph

GRAPH = "https://graph.microsoft.com/v1.0"


def _headers(extra=None):
    cabeceras = {"Authorization": f"Bearer {obtener_token_graph()}"}
    if extra:
        cabeceras.update(extra)
    return cabeceras


def _nombre_archivo(folio):
    limpio = re.sub(r'[\\/:*?"<>|]', "_", str(folio))
    return f"Oficio_{limpio}.docx"


def crear_y_abrir_word_en_nube(folio, asunto):
    """Sube un .docx EN BLANCO al OneDrive del usuario y lo abre en Word Online.

    Va en blanco a propósito: el folio y el asunto ya los imprime el PDF, y así
    todo el texto que se sincronice después es solo lo que redactó el usuario.
    Devuelve (file_id, web_url).
    """
    doc = Document()
    doc.core_properties.title = f"Oficio {folio}"
    doc.core_properties.subject = asunto
    buffer = io.BytesIO()
    doc.save(buffer)

    url = (
        f"{GRAPH}/me/drive/root:/{_nombre_archivo(folio)}:/content"
        "?@microsoft.graph.conflictBehavior=replace"
    )
    respuesta = requests.put(
        url,
        headers=_headers({
            "Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        }),
        data=buffer.getvalue(),
        timeout=60,
    )
    if respuesta.status_code not in (200, 201):
        raise RuntimeError(f"Error al crear el archivo en OneDrive: {respuesta.text}")

    datos = respuesta.json()
    file_id = datos.get("id")
    web_url = datos.get("webUrl")
    if web_url:
        webbrowser.open(web_url)
    return file_id, web_url


def obtener_texto_desde_onedrive(file_id):
    """Descarga el .docx actualizado desde OneDrive y devuelve su texto plano."""
    respuesta = requests.get(
        f"{GRAPH}/me/drive/items/{file_id}/content",
        headers=_headers(),
        timeout=60,
    )
    if respuesta.status_code != 200:
        raise RuntimeError(f"No se pudo sincronizar el contenido desde OneDrive: {respuesta.text}")

    doc = Document(io.BytesIO(respuesta.content))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]

    # Por si el usuario escribió dentro de tablas
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                partes.extend(p.text for p in celda.paragraphs if p.text.strip())

    return "\n".join(partes)


def descargar_pdf_desde_onedrive(file_id, ruta_destino_pdf):
    """Convierte y descarga el archivo editado desde OneDrive directamente en formato PDF."""
    respuesta = requests.get(
        f"{GRAPH}/me/drive/items/{file_id}/content?format=pdf",
        headers=_headers(),
        timeout=120,
    )
    if respuesta.status_code != 200:
        raise RuntimeError(f"No se pudo exportar el PDF desde OneDrive: {respuesta.text}")
    with open(ruta_destino_pdf, "wb") as archivo:
        archivo.write(respuesta.content)