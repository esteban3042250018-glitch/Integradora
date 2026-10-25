from io import BytesIO
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# Permite ejecutar este archivo directamente desde model/ y localizar controller/.
CARPETA_CODIGO = Path(__file__).resolve().parents[1]
if str(CARPETA_CODIGO) not in sys.path:
    sys.path.insert(0, str(CARPETA_CODIGO))

from controller.database_conecction import Database_connection


CARPETA_CODIGO = Path(__file__).resolve().parents[1]
CARPETA_OFICIOS = CARPETA_CODIGO / "oficios"

ANCHO_PAGINA = 8.5
ALTO_PAGINA = 11.0

PREFIJO_FOLIO = "CVDR-DGO-"
DIGITOS_FOLIO = 4

MARGEN_SUPERIOR = 0.85
MARGEN_INFERIOR = 0.85
MARGEN_IZQUIERDO = 0.85
MARGEN_DERECHO = 0.85


def obtener_conexion():
    """Abre una conexión usando el conector existente del proyecto."""
    conexion = Database_connection().get_connection()
    if conexion is None:
        raise ConnectionError(
            "No fue posible conectar con PostgreSQL. Revisa controller/config.py."
        )
    return conexion


def reservar_folio(conexion):
    """
    Reserva un folio único para la generación del documento.

    Bloquea la fila única de configuracion_folios, avanza el contador y
    comprueba ambas tablas de oficios para tolerar reinicios del contador.
    El contador se confirma antes de generar el DOCX: si la generación falla,
    puede quedar un número sin usar, pero nunca se reutiliza por accidente.
    """
    with conexion.cursor() as cursor:
        cursor.execute(
            """
            SELECT siguiente_numero
            FROM personalizacion.configuracion_folios
            WHERE id_configuracion = 1
            FOR UPDATE
            """
        )
        fila = cursor.fetchone()

        if fila is None:
            raise RuntimeError(
                "No existe la configuración de folios. "
                "Inserta la fila id_configuracion = 1 en "
                "personalizacion.configuracion_folios."
            )

        numero = fila[0]

        while True:
            folio = f"{PREFIJO_FOLIO}{numero:0{DIGITOS_FOLIO}d}"

            cursor.execute(
                """
                SELECT EXISTS (
                    SELECT 1
                    FROM correspondencia.oficios_envio
                    WHERE folio = %s
                )
                OR EXISTS (
                    SELECT 1
                    FROM correspondencia.oficios_recibido
                    WHERE folio = %s
                )
                """,
                (folio, folio),
            )
            ya_existe = cursor.fetchone()[0]

            if not ya_existe:
                break

            numero += 1

        cursor.execute(
            """
            UPDATE personalizacion.configuracion_folios
            SET siguiente_numero = %s,
                fecha_actualizacion = CURRENT_TIMESTAMP
            WHERE id_configuracion = 1
            """,
            (numero + 1,),
        )

    # Confirma el avance y libera el bloqueo de la fila de configuración.
    conexion.commit()
    return folio


def obtener_imagen_membrete(conexion):
    """Devuelve los bytes de la única imagen activa del membrete."""
    with conexion.cursor() as cursor:
        cursor.execute(
            """
            SELECT nombre_archivo, tipo_mime, imagen
            FROM personalizacion.membrete
            WHERE estatus = TRUE
            ORDER BY fecha_actualizacion DESC, id_membrete DESC
            LIMIT 2
            """
        )
        filas = cursor.fetchall()

    if not filas:
        raise RuntimeError(
            "No hay un membrete activo en personalizacion.membrete. "
            "Usa la opción 2 del menú de administración para registrar uno."
        )

    if len(filas) > 1:
        raise RuntimeError(
            "Hay más de un membrete activo. Corrige los estatus para dejar "
            "solo un registro con estatus = TRUE."
        )

    nombre_archivo, tipo_mime, imagen = filas[0]

    if not imagen:
        raise RuntimeError(
            f"El membrete activo ({nombre_archivo}) no contiene datos de imagen."
        )

    return nombre_archivo, tipo_mime, bytes(imagen)


def agregar_imagen_fondo(parrafo, imagen_bytes, ancho, alto):
    """Inserta la imagen almacenada en BYTEA como fondo detrás del texto."""
    parrafo.paragraph_format.space_before = Pt(0)
    parrafo.paragraph_format.space_after = Pt(0)
    parrafo.paragraph_format.left_indent = Inches(0)
    parrafo.paragraph_format.right_indent = Inches(0)

    imagen_stream = BytesIO(imagen_bytes)
    inline = parrafo.add_run().add_picture(
        imagen_stream,
        width=Inches(ancho),
        height=Inches(alto),
    )._inline

    ancla = OxmlElement("wp:anchor")

    propiedades = {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "1",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }

    for clave, valor in propiedades.items():
        ancla.set(clave, valor)

    posicion_simple = OxmlElement("wp:simplePos")
    posicion_simple.set("x", "0")
    posicion_simple.set("y", "0")
    ancla.append(posicion_simple)

    posicion_horizontal = OxmlElement("wp:positionH")
    posicion_horizontal.set("relativeFrom", "page")
    desplazamiento_x = OxmlElement("wp:posOffset")
    desplazamiento_x.text = "0"
    posicion_horizontal.append(desplazamiento_x)
    ancla.append(posicion_horizontal)

    posicion_vertical = OxmlElement("wp:positionV")
    posicion_vertical.set("relativeFrom", "page")
    desplazamiento_y = OxmlElement("wp:posOffset")
    desplazamiento_y.text = "0"
    posicion_vertical.append(desplazamiento_y)
    ancla.append(posicion_vertical)

    extension = inline.find(qn("wp:extent"))
    ancla.append(extension)

    extension_efectos = OxmlElement("wp:effectExtent")
    for lado in ("l", "t", "r", "b"):
        extension_efectos.set(lado, "0")
    ancla.append(extension_efectos)

    ancla.append(OxmlElement("wp:wrapNone"))

    for etiqueta in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        elemento = inline.find(qn(etiqueta))
        if elemento is not None:
            ancla.append(elemento)

    inline.getparent().replace(inline, ancla)


def agregar_folio(parrafo, folio):
    parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT

    formato = parrafo.paragraph_format
    formato.space_before = Inches(2.65)
    formato.space_after = Inches(0)
    formato.left_indent = Inches(0.65)
    formato.right_indent = Inches(0)
    formato.line_spacing = 1

    texto = parrafo.add_run(f"FOLIO: {folio}")
    texto.font.name = "Arial"
    texto.font.size = Pt(10)
    texto.bold = True




def crear_documento():

    """
    Obtiene el folio y la imagen activa de PostgreSQL y genera el DOCX.
    No usa icon.png ni un folio fijo.
    """
    conexion = obtener_conexion()

    try:
        # Primero reservamos el folio; después leemos la imagen activa.
        # La conexión se cierra en finally aunque ocurra un error.
        folio = reservar_folio(conexion)
        nombre_imagen, tipo_mime, imagen_bytes = obtener_imagen_membrete(conexion)
    finally:
        conexion.close()

    documento = Document()
    seccion = documento.sections[0]

    seccion.page_width = Inches(ANCHO_PAGINA)
    seccion.page_height = Inches(ALTO_PAGINA)

    seccion.top_margin = Inches(MARGEN_SUPERIOR)
    seccion.bottom_margin = Inches(MARGEN_INFERIOR)
    seccion.left_margin = Inches(MARGEN_IZQUIERDO)
    seccion.right_margin = Inches(MARGEN_DERECHO)

    seccion.header_distance = Inches(0.1)
    seccion.footer_distance = Inches(0)

    encabezado = seccion.header
    encabezado.is_linked_to_previous = False

    # El folio y el fondo se colocan en el encabezado para mantener el fondo
    # detrás del contenido de la página.
    parrafo_folio = encabezado.paragraphs[0]
    parrafo_folio.paragraph_format.space_before = Inches(1.65)
    agregar_folio(parrafo_folio, folio)

    parrafo_imagen = encabezado.add_paragraph()
    agregar_imagen_fondo(
        parrafo_imagen,
        imagen_bytes,
        ANCHO_PAGINA,
        ALTO_PAGINA,
    )

    cuerpo = documento.add_paragraph()
    cuerpo.paragraph_format.space_before = Pt(0)
    cuerpo.paragraph_format.space_after = Pt(0)

    # Crear la carpeta de destino si todavía no existe.
    CARPETA_OFICIOS.mkdir(parents=True, exist_ok=True)

    # El nombre del archivo usa el número del folio reservado:
    # CVDR-DGO-0001 -> Folio1.docx; CVDR-DGO-0012 -> Folio12.docx.
    numero_folio = int(folio.removeprefix(PREFIJO_FOLIO))
    ruta_salida = CARPETA_OFICIOS / f"Folio{numero_folio}.docx"
    documento.save(ruta_salida)
    print(f"Documento creado: {ruta_salida}")
    print(f"Folio reservado: {folio}")
    print(f"Membrete utilizado: {nombre_imagen} ({tipo_mime})")



if __name__ == "__main__":
    try:
        crear_documento()
    except Exception as error:
        print(f"No se pudo generar el membrete: {error}")
